from __future__ import annotations

import contextlib
import json
import threading
import time
from collections.abc import Callable, Iterator
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

import numpy as np
import pytest
from openpyxl import load_workbook

from app.core.context import DEFAULT_WORKBOOK_NAME, ApplicationContext
from app.core.lifecycle import LifecycleError, LifecycleService, MonitoringState, StartRequest
from app.core.settings import Settings
from app.domain.event_machine import SourceHealth
from app.domain.mutations import WorkbookMutation
from app.persistence.schema import WorkbookFailure
from app.persistence.writer import CommitReceipt
from app.services.orchestrator import (
    BusinessEventService,
    MasterData,
    OrderedAnalyzer,
    build_audit_mutation,
)
from app.video.evidence import EvidenceOutcome, EvidenceStatus
from app.video.sources import ReadKind, SourceFailure, SourceRead
from app.vision.detector import (
    Detection,
    DetectionBatch,
    DetectionContext,
    DetectionFailure,
    FrameEnvelope,
)
from tests.persistence.test_workbook import base_rows, write_workbook
from tools.create_mock_workbook import mock_rows

NOW = datetime(2026, 10, 8, 16, 0, tzinfo=UTC)
SESSION = UUID(int=0x5E55)
FULL_FRAME = np.zeros((720, 1280, 3), dtype=np.uint8)
WORKER_THREADS = {"pick-zone-capture", "pick-zone-analyzer", "pick-zone-preview"}
# shelf side is x < 640 px; the counting band is +/- 25.6 px; PICK confirms at x=750
PICK_PATH = [400] * 3 + list(range(450, 901, 50)) + [900] * 30
PICK_SEQUENCE = 9
COMPLETED_SEQUENCE = 29


# --------------------------------------------------------------------------- fakes


class RecordingWriter:
    """Stands in for WorkbookWriter inside the real WorkbookWriterWorker."""

    def __init__(self, *, fail: str | None = None, delay_s: float = 0.0) -> None:
        self.fail = fail
        self.delay_s = delay_s
        self.mutations: list[WorkbookMutation] = []
        self.lock = threading.Lock()

    def commit(self, mutation: WorkbookMutation) -> CommitReceipt:
        if self.delay_s:
            time.sleep(self.delay_s)
        if self.fail:
            raise WorkbookFailure(self.fail, "injected")
        with self.lock:
            self.mutations.append(mutation)
        version = mutation.expected_workbook_version
        return CommitReceipt(
            mutation_id=str(mutation.mutation_id),
            idempotency_key=mutation.idempotency_key,
            event_id="",
            mutation_hash=mutation.canonical_effects_hash,
            committed_at=NOW.isoformat(),
            workbook_version_before=version,
            workbook_version_after=version + 1,
        )

    def sheets(self, sheet: str) -> list[dict[str, object]]:
        with self.lock:
            return [
                row
                for mutation in self.mutations
                for row in (*mutation.upserts.get(sheet, ()), *mutation.appends.get(sheet, ()))
            ]


class FakeSource:
    def __init__(
        self,
        value: int = 1,
        *,
        frames: int | None = None,
        interval_s: float = 0.005,
        ts_step_ms: int = 33,
        fail_open: str | None = None,
        image: np.ndarray | None = None,
    ) -> None:
        self.value = value
        self.frames = frames
        self.interval_s = interval_s
        self.ts_step_ms = ts_step_ms
        self.fail_open = fail_open
        self.image = image if image is not None else np.full((72, 128, 3), value, np.uint8)
        self.continuity_segment = 0
        self.opened = False
        self.released = False
        self.paused = False
        self.reads = 0

    def open(self) -> None:
        if self.fail_open:
            raise SourceFailure(self.fail_open, "fake")
        self.opened = True

    def read(self) -> SourceRead:
        if self.released:
            raise SourceFailure("SOURCE_NOT_OPEN", "released")
        time.sleep(self.interval_s)
        if self.paused:
            return SourceRead(ReadKind.PAUSED)
        if self.frames is not None and self.reads >= self.frames:
            return SourceRead(ReadKind.ENDED)
        index = self.reads
        self.reads += 1
        return SourceRead(ReadKind.FRAME, self.image, index * self.ts_step_ms)

    def pause(self) -> None:
        self.paused = True

    def resume(self) -> None:
        self.paused = False

    def release(self) -> None:
        self.released = True


def box(x: float, confidence: float = 0.999) -> Detection:
    return Detection("d", 0, "box", confidence, (x - 50, 400, x + 50, 500))


class RecordingDetector:
    def __init__(
        self,
        *,
        delay_s: float = 0.0,
        positions: Callable[[int], float | None] | None = None,
        block: threading.Event | None = None,
    ) -> None:
        self.delay_s = delay_s
        self.positions = positions
        self.block = block
        self.calls: list[tuple[UUID, int, int]] = []

    def detect(self, frame: FrameEnvelope, context: DetectionContext) -> DetectionBatch:
        index = len(self.calls)
        self.calls.append((frame.session_id, frame.sequence, int(frame.image_bgr[0, 0, 0])))
        if self.block is not None:
            self.block.wait(10)
        if self.delay_s:
            time.sleep(self.delay_s)
        position = self.positions(index) if self.positions else None
        return DetectionBatch(
            session_id=frame.session_id,
            continuity_segment=frame.continuity_segment,
            frame_sequence=frame.sequence,
            source_timestamp_ms=frame.source_timestamp_ms,
            started_monotonic_ns=1,
            completed_monotonic_ns=2,
            model_id="fake",
            detections=() if position is None else (box(position),),
        )


class ScriptedDetector:
    """Returns one box per frame sequence following ``PICK_PATH``."""

    def __init__(self, path: list[int] | None = None, fail_at: int | None = None) -> None:
        self.path = path or PICK_PATH
        self.fail_at = fail_at

    def detect(self, frame: FrameEnvelope, context: DetectionContext) -> DetectionBatch:
        if frame.sequence == self.fail_at:
            raise DetectionFailure("INFERENCE_FAILED", "boom", frame.sequence)
        position = self.path[frame.sequence] if frame.sequence < len(self.path) else None
        return DetectionBatch(
            session_id=frame.session_id,
            continuity_segment=frame.continuity_segment,
            frame_sequence=frame.sequence,
            source_timestamp_ms=frame.source_timestamp_ms,
            started_monotonic_ns=1,
            completed_monotonic_ns=2,
            model_id="scripted",
            detections=() if position is None else (box(position),),
        )


class FakeEvidence:
    def __init__(self) -> None:
        self.begun: list[tuple[UUID, int]] = []
        self.completed: list[tuple[UUID, int]] = []
        self.abandoned: list[tuple[UUID, str]] = []
        self.frames = 0
        self.source_ended = False

    def add_frame(self, frame: FrameEnvelope) -> bool:
        self.frames += 1
        return True

    def begin_event(self, event_id: UUID, first_action_source_ms: int) -> None:
        self.begun.append((event_id, first_action_source_ms))

    def complete_event(self, event_id: UUID, ended_source_ms: int) -> None:
        self.completed.append((event_id, ended_source_ms))

    def abandon_event(self, event_id: UUID, reason: str) -> EvidenceOutcome:
        self.abandoned.append((event_id, reason))
        return EvidenceOutcome(event_id, EvidenceStatus.UNAVAILABLE, None, None, None, 0, reason)

    def mark_source_ended(self) -> None:
        self.source_ended = True

    def finalize_ready(self) -> tuple[EvidenceOutcome, ...]:
        return ()

    def shutdown(self, timeout_s: float | None = None) -> tuple[EvidenceOutcome, ...]:
        return ()


def available(event_id: UUID) -> EvidenceOutcome:
    return EvidenceOutcome(
        event_id, EvidenceStatus.AVAILABLE, f"evidence/s/{event_id}.mp4", 0, 9_000, 90, None
    )


# --------------------------------------------------------------------------- builders


def rows_with(**changes: Callable[[dict[str, list[dict[str, object]]]], None]):
    rows = base_rows()
    rows["Tasks"][0]["expected_quantity"] = 1
    rows["SourceProfiles"].append(
        {
            **rows["SourceProfiles"][0],
            "source_profile_id": "source-2",
            "name": "Camera",
            "source_type": "camera",
            "file_path": "",
            "camera_index": 0,
            "active": False,
        }
    )
    rows["ZoneProfiles"].append(
        {**rows["ZoneProfiles"][0], "zone_profile_id": "zone-2", "source_profile_id": "source-2"}
    )
    for change in changes.values():
        change(rows)
    return rows


class Harness:
    def __init__(self, tmp_path: Path) -> None:
        self.tmp_path = tmp_path
        self.contexts: list[ApplicationContext] = []
        self.services: list[LifecycleService] = []

    def context(
        self,
        *,
        rows: dict[str, list[dict[str, object]]] | None = None,
        write: bool = True,
        writer: RecordingWriter | str | None = "recording",
        sources: dict[str, list[FakeSource]] | None = None,
        detector: object | None = None,
        evidence_factory: Callable[[UUID], object | None] = lambda _session: None,
        restorer: Callable[[str], None] | None = None,
        which: Callable[[str], str | None] = lambda _name: "/usr/bin/ffmpeg",
        detector_check: Callable[[], str | None] = lambda: None,
    ) -> ApplicationContext:
        settings = Settings(project_dir=self.tmp_path)
        settings.prepare_directories()
        assert settings.data_dir is not None
        if write:
            write_workbook(settings.data_dir / DEFAULT_WORKBOOK_NAME, rows or rows_with())
        queues = {key: list(value) for key, value in (sources or {}).items()}

        def source_factory(profile: dict[str, object], override: str | None) -> FakeSource:
            return queues[str(profile["source_profile_id"])].pop(0)

        self.writer = RecordingWriter() if writer == "recording" else writer
        context = ApplicationContext.create(
            settings,
            writer=self.writer,
            source_factory=source_factory,
            detector_factory=lambda _profile, _classes: detector or RecordingDetector(),
            evidence_factory=evidence_factory,
            scenario_restorer=restorer,
            detector_check=detector_check,
            which=which,
        )
        context.start()
        self.contexts.append(context)
        return context

    def lifecycle(self, context: ApplicationContext, **kwargs: object) -> LifecycleService:
        service = LifecycleService(context, **kwargs)  # type: ignore[arg-type]
        self.services.append(service)
        return service

    def close(self) -> None:
        for service in self.services:
            with contextlib.suppress(LifecycleError):
                service.shutdown(timeout_s=5)
        for context in self.contexts:
            context.close(timeout_s=5)


@pytest.fixture
def harness(tmp_path: Path) -> Iterator[Harness]:
    instance = Harness(tmp_path)
    yield instance
    instance.close()


def request(source: str = "source-1", zone: str = "zone-1", **kwargs: object) -> StartRequest:
    return StartRequest(source_profile_id=source, zone_profile_id=zone, task_id="task-1", **kwargs)


def wait_until(predicate: Callable[[], bool], timeout: float = 3.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return True
        time.sleep(0.01)
    return predicate()


def live_worker_names() -> set[str]:
    return {thread.name for thread in threading.enumerate() if thread.name.startswith("pick-zone-")}


def error_of(call: Callable[[], object]) -> LifecycleError:
    with pytest.raises(LifecycleError) as caught:
        call()
    return caught.value


# --------------------------------------------------------------------------- pipeline helpers


class Pipeline:
    def __init__(
        self,
        context: ApplicationContext,
        *,
        evidence: FakeEvidence | None = None,
        detector: object | None = None,
    ) -> None:
        snapshot = context.load_snapshot()
        self.context = context
        self.master = MasterData.from_snapshot(snapshot, "task-1")
        self.zone = snapshot.indexes.active_zone_by_source["source-1"]
        self.business = BusinessEventService(
            context.persistence,
            context.runtime,
            self.master,
            SESSION,
            reload=context.load_snapshot,
            now=lambda: NOW,
        )
        self.analyzer = OrderedAnalyzer(
            session_id=SESSION,
            zone=self.zone,
            event_context=self.master.event_context(self.zone, SESSION, "test-config"),
            detector=detector or ScriptedDetector(),
            runtime=context.runtime,
            business=self.business,
            allowed_class_names=frozenset({"box"}),
            evidence=evidence,
            wall_now=lambda: NOW,
        )
        context.runtime.reset_session(SESSION)

    def feed(self, sequences: range, *, offset_ms: int = 0) -> None:
        for sequence in sequences:
            self.analyzer.analyze(frame(sequence, offset_ms=offset_ms))


def frame(sequence: int, *, offset_ms: int = 0) -> FrameEnvelope:
    return FrameEnvelope(
        session_id=SESSION,
        continuity_segment=0,
        sequence=sequence,
        source_timestamp_ms=sequence * 250 + offset_ms,
        captured_monotonic_ns=sequence,
        image_bgr=FULL_FRAME,
    )


# --------------------------------------------------------------------------- readiness


def test_readiness_reports_every_missing_runtime_prerequisite(harness: Harness) -> None:
    context = harness.context(
        write=False,
        which=lambda _name: None,
        detector_check=lambda: "DETECTOR_UNAVAILABLE",
    )

    report = context.readiness()

    assert report.status == "degraded"
    assert {"WORKBOOK_NOT_FOUND", "EVIDENCE_UNAVAILABLE", "DETECTOR_UNAVAILABLE"} <= set(
        report.problems
    )


def test_readiness_is_ready_when_workbook_storage_detector_and_ffmpeg_are_available(
    harness: Harness,
) -> None:
    context = harness.context()

    report = context.readiness()

    assert report.status == "ready"
    assert report.problems == ()
    assert context.workbook_path.name == DEFAULT_WORKBOOK_NAME
    assert context.persistence.writable


# --------------------------------------------------------------------------- start validation


def test_start_identifies_every_missing_prerequisite_and_starts_nothing(harness: Harness) -> None:
    def break_master_data(rows: dict[str, list[dict[str, object]]]) -> None:
        rows["Tasks"][0]["status"] = "completed"
        rows["Skus"][0]["active"] = False
        rows["Locations"][1]["active"] = False
        rows["Inventory"].pop(1)

    source = FakeSource()
    context = harness.context(
        rows=rows_with(change=break_master_data), sources={"source-1": [source]}
    )
    lifecycle = harness.lifecycle(context)

    error = error_of(lambda: lifecycle.start(request(zone="zone-missing")))

    assert error.code == "ZONE_PROFILE_NOT_FOUND"
    assert error.status == 404
    missing = {(item["kind"], item["id"], item["reason"]) for item in error.details["missing"]}
    assert missing == {
        ("zone_profile", "zone-missing", "not_found"),
        ("task", "task-1", "not_open"),
        ("sku", "sku-box", "inactive"),
        ("destination_location", "STAGING", "inactive"),
        ("destination_inventory", "STAGING/sku-box", "not_found"),
    }
    assert lifecycle.state is MonitoringState.STOPPED
    assert not source.opened
    assert not (live_worker_names() & WORKER_THREADS)


def mock_request(task_id: str) -> StartRequest:
    return StartRequest(
        source_profile_id="SRC-FILE-01", zone_profile_id="ZONE-A-03-01", task_id=task_id
    )


def test_mock_task_at_the_zone_location_starts_monitoring(harness: Harness) -> None:
    source = FakeSource()
    context = harness.context(rows=mock_rows(), sources={"SRC-FILE-01": [source]})
    lifecycle = harness.lifecycle(context)

    lifecycle.start(mock_request("PICK-1001"))

    assert lifecycle.state is MonitoringState.RUNNING
    snapshot = context.runtime.snapshot()
    assert snapshot.task is not None and snapshot.task.task_id == "PICK-1001"
    assert snapshot.inventory is not None and snapshot.inventory.location_id == "A-03-01"
    lifecycle.stop()


def test_mock_task_at_another_location_is_a_profile_task_mismatch(harness: Harness) -> None:
    source = FakeSource()
    context = harness.context(rows=mock_rows(), sources={"SRC-FILE-01": [source]})
    lifecycle = harness.lifecycle(context)

    error = error_of(lambda: lifecycle.start(mock_request("PICK-1002")))

    assert (error.code, error.status) == ("PROFILE_TASK_MISMATCH", 409)
    assert error.details["missing"] == [
        {"kind": "zone_location", "id": "A-03-01", "reason": "task_mismatch"}
    ]
    assert lifecycle.state is MonitoringState.STOPPED
    assert not source.opened


def test_mismatch_ranks_after_not_found_codes_and_before_validation_error(
    harness: Harness,
) -> None:
    def drop_rice_staging(rows: dict[str, list[dict[str, object]]]) -> None:
        rows["Inventory"] = [
            item
            for item in rows["Inventory"]
            if (item["location_id"], item["sku_id"]) != ("STAGING-01", "SKU-RICE-10KG")
        ]

    rows = mock_rows()
    drop_rice_staging(rows)
    context = harness.context(rows=rows)
    lifecycle = harness.lifecycle(context)

    assert error_of(lambda: lifecycle.start(mock_request("PICK-1002"))).code == (
        "PROFILE_TASK_MISMATCH"
    )
    assert error_of(lambda: lifecycle.start(mock_request("PICK-404"))).code == "TASK_NOT_FOUND"


def test_start_uses_contract_codes_in_precedence_order(harness: Harness) -> None:
    context = harness.context()
    lifecycle = harness.lifecycle(context)

    assert error_of(lambda: lifecycle.start(request(source="nope"))).code == (
        "SOURCE_PROFILE_NOT_FOUND"
    )
    # a zone that belongs to another source is not a valid zone for this source
    assert error_of(lambda: lifecycle.start(request(zone="zone-2"))).code == (
        "ZONE_PROFILE_NOT_FOUND"
    )
    missing_task = StartRequest("source-1", "zone-1", "task-nope")
    assert error_of(lambda: lifecycle.start(missing_task)).code == "TASK_NOT_FOUND"


def test_start_rejects_invalid_analysis_rate_and_misplaced_network_override(
    harness: Harness,
) -> None:
    lifecycle = harness.lifecycle(harness.context())

    fps_error = error_of(lambda: lifecycle.start(request(analysis_fps=6)))
    override_error = error_of(
        lambda: lifecycle.start(request(network_url_override="rtsp://user:secret@cam/1"))
    )

    assert (fps_error.code, fps_error.status) == ("VALIDATION_ERROR", 400)
    assert override_error.code == "VALIDATION_ERROR"
    assert "secret" not in str(override_error) and "secret" not in json.dumps(
        override_error.details
    )


def test_source_open_failure_reports_422_and_starts_no_workers(harness: Harness) -> None:
    source = FakeSource(fail_open="DECODE_FAILED")
    context = harness.context(sources={"source-1": [source]})
    lifecycle = harness.lifecycle(context)

    error = error_of(lambda: lifecycle.start(request()))

    assert (error.code, error.status) == ("SOURCE_OPEN_FAILED", 422)
    assert error.details["source_error"] == "DECODE_FAILED"
    assert lifecycle.state is MonitoringState.STOPPED
    assert context.runtime.snapshot().monitoring.state == "stopped"
    assert not (live_worker_names() & WORKER_THREADS)


def test_start_is_blocked_when_the_workbook_cannot_accept_decisions(harness: Harness) -> None:
    context = harness.context(sources={"source-1": [FakeSource()]})
    context.workbook_path.unlink()
    lifecycle = harness.lifecycle(context)

    error = error_of(lambda: lifecycle.start(request()))

    assert (error.code, error.status) == ("PERSISTENCE_BLOCKED", 409)
    assert lifecycle.state is MonitoringState.STOPPED


# --------------------------------------------------------------------------- worker graph


def test_start_builds_worker_graph_and_publishes_running_state(harness: Harness) -> None:
    evidence = FakeEvidence()
    source = FakeSource()
    detector = RecordingDetector()
    context = harness.context(
        sources={"source-1": [source]},
        detector=detector,
        evidence_factory=lambda _session: evidence,
    )
    lifecycle = harness.lifecycle(context)

    result = lifecycle.start(request(analysis_fps=5))

    assert result.state == "starting"
    assert lifecycle.state is MonitoringState.RUNNING
    assert source.opened
    assert WORKER_THREADS | {"pick-zone-evidence"} <= lifecycle.worker_names()
    assert wait_until(lambda: len(detector.calls) >= 2 and evidence.frames >= 2)
    snapshot = context.runtime.snapshot()
    assert snapshot.monitoring.state == "running"
    assert snapshot.monitoring.session_id == result.session_id
    assert snapshot.monitoring.source_profile_id == "source-1"
    assert snapshot.monitoring.zone_profile_id == "zone-1"
    assert snapshot.monitoring.task_id == "task-1"
    assert snapshot.source.state == "connected"
    assert snapshot.task is not None and snapshot.task.task_id == "task-1"
    assert snapshot.inventory is not None and snapshot.inventory.quantity == 10
    assert wait_until(lambda: context.runtime.preview() is not None)
    assert wait_until(
        lambda: any(row.get("status") == "running" for row in harness.writer.sheets("Sessions"))
    )

    assert error_of(lambda: lifecycle.start(request())).code == "MONITORING_ALREADY_ACTIVE"
    lifecycle.stop()
    assert not (live_worker_names() & (WORKER_THREADS | {"pick-zone-evidence"}))


def test_missing_evidence_encoder_degrades_with_visible_alert(harness: Harness) -> None:
    context = harness.context(sources={"source-1": [FakeSource()]})
    lifecycle = harness.lifecycle(context)

    lifecycle.start(request())

    assert context.runtime.has_alert("EVIDENCE_UNAVAILABLE")
    assert "pick-zone-evidence" not in lifecycle.worker_names()


# --------------------------------------------------------------------------- transitions


def test_pause_resume_and_invalid_transitions(harness: Harness) -> None:
    source = FakeSource()
    detector = RecordingDetector()
    context = harness.context(sources={"source-1": [source]}, detector=detector)
    lifecycle = harness.lifecycle(context)

    assert error_of(lifecycle.pause).code == "INVALID_MONITORING_STATE"
    assert error_of(lifecycle.resume).code == "INVALID_MONITORING_STATE"
    assert error_of(lifecycle.stop).code == "INVALID_MONITORING_STATE"

    lifecycle.start(request(analysis_fps=5))
    assert wait_until(lambda: len(detector.calls) >= 1)
    assert lifecycle.pause() is MonitoringState.PAUSED
    assert source.paused
    assert context.runtime.snapshot().monitoring.state == "paused"
    assert error_of(lifecycle.pause).status == 409
    calls_when_paused = len(detector.calls)
    time.sleep(0.5)
    assert len(detector.calls) <= calls_when_paused + 1  # at most the in-flight frame

    assert lifecycle.resume() is MonitoringState.RUNNING
    assert not source.paused
    assert error_of(lifecycle.resume).code == "INVALID_MONITORING_STATE"
    assert wait_until(lambda: len(detector.calls) > calls_when_paused + 1)

    summary = lifecycle.stop()
    assert summary.state == "stopped"
    assert lifecycle.state is MonitoringState.STOPPED
    assert source.released
    assert context.runtime.snapshot().monitoring.state == "stopped"
    assert any(row.get("status") == "stopped" for row in harness.writer.sheets("Sessions"))


def test_resume_reports_source_unavailable_after_the_source_ended(harness: Harness) -> None:
    context = harness.context(sources={"source-1": [FakeSource(frames=3)]})
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request())
    assert wait_until(lambda: context.runtime.snapshot().source.state == "ended")

    lifecycle.pause()
    error = error_of(lifecycle.resume)

    assert (error.code, error.status) == ("SOURCE_UNAVAILABLE", 422)
    assert lifecycle.state is MonitoringState.PAUSED


def test_stop_requires_decision_for_active_event_and_discard_changes_no_inventory(
    harness: Harness,
) -> None:
    path = [400.0] * 3 + [float(x) for x in range(450, 901, 50)]
    source = FakeSource(interval_s=0.01, ts_step_ms=10, image=FULL_FRAME)
    detector = RecordingDetector(positions=lambda index: path[min(index, len(path) - 1)])
    context = harness.context(sources={"source-1": [source]}, detector=detector)
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request(analysis_fps=5))
    assert wait_until(lambda: context.runtime.snapshot().event.state == "active", timeout=6)
    event_id = context.runtime.snapshot().event.event_id

    error = error_of(lambda: lifecycle.stop(discard_in_progress=False))
    assert error.code == "ACTIVE_EVENT_REQUIRES_DECISION"
    assert error.status == 409
    assert lifecycle.state is MonitoringState.RUNNING

    summary = lifecycle.stop(discard_in_progress=True)

    assert summary.discarded_event_id == event_id
    assert lifecycle.state is MonitoringState.STOPPED
    audits = harness.writer.sheets("AuditLog")
    assert any(
        row["action"] == "event_discarded" and row["entity_id"] == str(event_id) for row in audits
    )
    assert harness.writer.sheets("Inventory") == []
    assert harness.writer.sheets("Events") == []


# --------------------------------------------------------------------------- reset


def test_reset_rules_and_audited_restore(harness: Harness) -> None:
    restored: list[str] = []

    def restorer(scenario_id: str) -> None:
        if scenario_id == "broken":
            raise WorkbookFailure("TEMP_SAVE_FAILED", "disk full")
        if scenario_id != "basic-pick":
            raise KeyError(scenario_id)
        restored.append(scenario_id)

    context = harness.context(sources={"source-1": [FakeSource()]}, restorer=restorer)
    lifecycle = harness.lifecycle(context)

    assert error_of(lambda: lifecycle.reset("basic-pick", confirm=False)).code == (
        "VALIDATION_ERROR"
    )
    not_found = error_of(lambda: lifecycle.reset("missing", confirm=True))
    assert (not_found.code, not_found.status) == ("SCENARIO_NOT_FOUND", 404)
    blocked = error_of(lambda: lifecycle.reset("broken", confirm=True))
    assert (blocked.code, blocked.status) == ("PERSISTENCE_BLOCKED", 503)

    result = lifecycle.reset("basic-pick", confirm=True)

    assert restored == ["basic-pick"]
    assert result.scenario_id == "basic-pick"
    assert [(item.location_id, item.quantity) for item in result.inventory] == [
        ("A-03-01", 10),
        ("STAGING", 0),
    ]
    audits = harness.writer.sheets("AuditLog")
    assert [row["entity_id"] for row in audits if row["action"] == "scenario_reset"] == [
        "basic-pick"
    ]

    lifecycle.start(request())
    active = error_of(lambda: lifecycle.reset("basic-pick", confirm=True))
    assert (active.code, active.status) == ("MONITORING_ACTIVE", 409)
    assert restored == ["basic-pick"]


# --------------------------------------------------------------------------- source switching


def test_switching_sources_releases_previous_source_and_isolates_sessions(
    harness: Harness,
) -> None:
    first = FakeSource(value=11)
    second = FakeSource(value=22)
    detector = RecordingDetector()
    context = harness.context(
        sources={"source-1": [first], "source-2": [second]}, detector=detector
    )
    lifecycle = harness.lifecycle(context)

    first_session = lifecycle.start(request(analysis_fps=5)).session_id
    assert wait_until(lambda: len(detector.calls) >= 2)
    # switching while active is refused until the session is stopped
    assert error_of(lambda: lifecycle.start(request("source-2", "zone-2"))).code == (
        "MONITORING_ALREADY_ACTIVE"
    )
    lifecycle.stop()
    assert first.released
    calls_before_switch = len(detector.calls)

    second_session = lifecycle.start(request("source-2", "zone-2", analysis_fps=5)).session_id
    assert wait_until(lambda: len(detector.calls) >= calls_before_switch + 3)
    assert wait_until(lambda: context.runtime.latest_frame() is not None)

    assert second_session != first_session
    new_calls = detector.calls[calls_before_switch:]
    assert {(session, value) for session, _, value in new_calls} == {(second_session, 22)}
    latest = context.runtime.latest_frame()
    assert latest is not None and latest.session_id == second_session
    assert int(latest.image_bgr[0, 0, 0]) == 22
    assert context.runtime.snapshot().monitoring.source_profile_id == "source-2"
    assert not second.released


# --------------------------------------------------------------------------- backpressure


def test_fast_source_with_slow_detector_keeps_queues_bounded_and_controls_responsive(
    harness: Harness,
) -> None:
    source = FakeSource(interval_s=0.001)
    detector = RecordingDetector(delay_s=0.1)
    context = harness.context(sources={"source-1": [source]}, detector=detector)
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request(analysis_fps=5))

    depths: list[tuple[int, int]] = []
    deadline = time.monotonic() + 1.2
    while time.monotonic() < deadline:
        metrics = context.runtime.snapshot().metrics
        depths.append((metrics.analysis_queue_depth, metrics.persistence_queue_depth))
        time.sleep(0.02)

    assert max(analysis for analysis, _ in depths) <= 1
    assert max(persistence for _, persistence in depths) <= 100
    metrics = context.runtime.snapshot().metrics
    assert metrics.captured_frames >= 100
    assert metrics.analyzed_frames <= 10
    assert metrics.dropped_frames >= metrics.captured_frames // 2
    preview = context.runtime.preview()
    assert preview is not None and preview.frame_sequence > 50  # preview ignores inference
    assert metrics.capture_fps > metrics.analysis_fps

    started = time.monotonic()
    lifecycle.pause()
    assert time.monotonic() - started < 0.5


def test_persistence_failure_blocks_auto_approval_and_next_start(harness: Harness) -> None:
    context = harness.context(
        sources={"source-1": [FakeSource(), FakeSource()]},
        writer=RecordingWriter(fail="TEMP_SAVE_FAILED"),
    )
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request())

    assert wait_until(lambda: context.runtime.snapshot().persistence.state == "blocked")
    assert context.runtime.has_alert("PERSISTENCE_BLOCKED")
    assert not context.persistence.writable
    assert context.readiness().status == "degraded"
    lifecycle.stop()

    error = error_of(lambda: lifecycle.start(request()))
    assert (error.code, error.status) == ("PERSISTENCE_BLOCKED", 409)


# --------------------------------------------------------------------------- shutdown


def test_shutdown_stops_workers_releases_source_and_flushes_within_five_seconds(
    harness: Harness,
) -> None:
    source = FakeSource()
    context = harness.context(
        sources={"source-1": [source]}, detector=RecordingDetector(delay_s=0.3)
    )
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request(analysis_fps=5))
    time.sleep(0.4)

    started = time.monotonic()
    lifecycle.shutdown(timeout_s=5)
    elapsed = time.monotonic() - started

    assert elapsed < 5
    assert source.released
    assert lifecycle.state is MonitoringState.STOPPED
    assert not (live_worker_names() & WORKER_THREADS)
    assert "pick-zone-workbook-writer" not in live_worker_names()
    assert any(row.get("status") == "stopped" for row in harness.writer.sheets("Sessions"))
    assert error_of(lambda: lifecycle.start(request())).code == "INVALID_MONITORING_STATE"


def test_hung_inference_reports_shutdown_timeout(harness: Harness) -> None:
    block = threading.Event()
    detector = RecordingDetector(block=block)
    context = harness.context(sources={"source-1": [FakeSource()]}, detector=detector)
    lifecycle = harness.lifecycle(context)
    lifecycle.start(request(analysis_fps=5))
    assert wait_until(lambda: len(detector.calls) >= 1)

    started = time.monotonic()
    error = error_of(lambda: lifecycle.stop(discard_in_progress=True, timeout_s=0.5))
    elapsed = time.monotonic() - started

    assert (error.code, error.status) == ("SHUTDOWN_TIMEOUT", 503)
    assert elapsed < 2
    assert lifecycle.state is MonitoringState.FAILED
    block.set()
    assert wait_until(lambda: not (live_worker_names() & WORKER_THREADS))


# --------------------------------------------------------------------------- ordered analyzer


def test_completed_event_waits_for_evidence_then_auto_approves(harness: Harness) -> None:
    evidence = FakeEvidence()
    context = harness.context()
    pipeline = Pipeline(context, evidence=evidence)

    pipeline.feed(range(PICK_SEQUENCE + 1))
    snapshot = context.runtime.snapshot()
    assert snapshot.event.state == "active"
    assert snapshot.event.pick_count == 1
    assert snapshot.analysis.actions and snapshot.analysis.actions[0].action_type == "pick"
    event_id = snapshot.event.event_id
    assert evidence.begun == [(event_id, PICK_SEQUENCE * 250)]

    pipeline.feed(range(PICK_SEQUENCE + 1, len(PICK_PATH)))
    assert evidence.completed == [(event_id, COMPLETED_SEQUENCE * 250)]
    assert pipeline.business.decisions == []  # waiting for the clip
    assert context.runtime.snapshot().event.state == "completed"

    pipeline.business.resolve_evidence(available(event_id))
    assert context.persistence.drain(timeout_s=2)

    decision = pipeline.business.decisions[-1]
    assert (decision.event_id, decision.decision, decision.persisted) == (
        event_id,
        "auto_approved",
        True,
    )
    [mutation] = harness.writer.mutations
    assert mutation.expected_workbook_version == 0
    inventory = {row["location_id"]: row["quantity"] for row in mutation.upserts["Inventory"]}
    assert inventory == {"A-03-01": 9, "STAGING": 1}
    assert mutation.upserts["Events"][0]["evidence_path"] == f"evidence/s/{event_id}.mp4"
    assert context.runtime.snapshot().event.state == "approved"
    assert context.persistence.workbook_version == 1


def test_review_required_event_is_persisted_to_the_real_workbook(harness: Harness) -> None:
    context = harness.context(writer=None)
    pipeline = Pipeline(context)

    pipeline.feed(range(len(PICK_PATH)))
    assert context.persistence.drain(timeout_s=5)

    decision = pipeline.business.decisions[-1]
    assert (decision.decision, decision.review_reason, decision.persisted) == (
        "review_required",
        "EVIDENCE_UNAVAILABLE",
        True,
    )
    workbook = load_workbook(context.workbook_path)

    def rows(sheet: str) -> list[dict[str, object]]:
        values = list(workbook[sheet].iter_rows(values_only=True))
        return [dict(zip(values[0], row, strict=True)) for row in values[1:]]

    [event] = rows("Events")
    assert event["event_id"] == str(decision.event_id)
    assert event["state"] == "review_required"
    assert event["decision"] == "pending"
    assert event["review_reason"] == "EVIDENCE_UNAVAILABLE"
    assert event["session_id"] == str(SESSION)
    assert event["task_id"] == "task-1"
    assert (event["zone_profile_id"], event["zone_version"]) == ("zone-1", 1)
    assert event["observed_quantity"] == 1
    assert "EVIDENCE_UNAVAILABLE" in json.loads(str(event["integrity_flags_json"]))
    [action] = rows("EventActions")
    assert action["action_type"] == "pick"
    assert json.loads(str(action["bbox_json"])) == [700.0, 400.0, 800.0, 500.0]
    assert any(row["action"] == "review_required" for row in rows("AuditLog"))
    assert {row["location_id"]: row["quantity"] for row in rows("Inventory")} == {
        "A-03-01": 10,
        "STAGING": 0,
    }
    assert context.runtime.snapshot().event.state == "review_required"
    assert context.runtime.snapshot().event.review_reason == "EVIDENCE_UNAVAILABLE"


def test_blocked_persistence_keeps_event_review_required_in_memory(harness: Harness) -> None:
    writer = RecordingWriter(fail="ATOMIC_REPLACE_FAILED")
    context = harness.context(writer=writer)
    evidence = FakeEvidence()
    pipeline = Pipeline(context, evidence=evidence)
    context.persistence.submit(
        lambda version: build_audit_mutation(
            version,
            idempotency_key="probe",
            action="probe",
            entity_type="test",
            entity_id="probe",
            timestamp=NOW,
        )
    )
    assert context.persistence.drain(timeout_s=2)
    assert context.persistence.state == "blocked"

    pipeline.feed(range(len(PICK_PATH)))
    event_id = evidence.completed[0][0]
    pipeline.business.resolve_evidence(available(event_id))

    decision = pipeline.business.decisions[-1]
    assert (decision.decision, decision.review_reason, decision.persisted) == (
        "review_required",
        "PERSISTENCE_BLOCKED",
        False,
    )
    assert event_id in pipeline.business.unpersisted_events
    assert writer.mutations == []
    snapshot = context.runtime.snapshot()
    assert snapshot.event.state == "review_required"
    assert snapshot.persistence.state == "blocked"
    assert context.runtime.has_alert("PERSISTENCE_BLOCKED")


def test_late_results_are_discarded_counted_and_flag_the_active_event(harness: Harness) -> None:
    context = harness.context()
    evidence = FakeEvidence()
    pipeline = Pipeline(context, evidence=evidence)
    pipeline.feed(range(PICK_SEQUENCE + 2))

    assert pipeline.analyzer.analyze(frame(PICK_SEQUENCE - 1)) is None

    assert context.runtime.snapshot().metrics.late_results == 1
    assert context.runtime.snapshot().analysis.last_applied_frame_sequence == PICK_SEQUENCE + 1
    assert "LATE_RESULT_DISCARDED" in pipeline.analyzer.active_event().integrity_flags
    pipeline.feed(range(PICK_SEQUENCE + 2, len(PICK_PATH)))
    pipeline.business.resolve_evidence(available(evidence.completed[0][0]))
    assert context.persistence.drain(timeout_s=2)
    assert pipeline.business.decisions[-1].review_reason == "LATE_RESULT_DISCARDED"


def test_paused_time_does_not_advance_event_timers(harness: Harness) -> None:
    context = harness.context()
    pipeline = Pipeline(context)
    pipeline.feed(range(PICK_SEQUENCE + 4))

    pipeline.analyzer.pause()
    pipeline.analyzer.resume()
    pipeline.feed(range(PICK_SEQUENCE + 4, PICK_SEQUENCE + 8), offset_ms=60_000)

    event = pipeline.analyzer.active_event()
    assert event.state.value in {"active", "settling"}
    assert event.review_reason is None


def test_source_interruption_sends_active_event_to_review(harness: Harness) -> None:
    context = harness.context()
    evidence = FakeEvidence()
    pipeline = Pipeline(context, evidence=evidence)
    pipeline.feed(range(PICK_SEQUENCE + 2))
    event_id = context.runtime.snapshot().event.event_id

    pipeline.analyzer.notify_source_health(SourceHealth.RECONNECTING)
    pipeline.analyzer.process_pending()

    assert evidence.completed and evidence.completed[0][0] == event_id
    pipeline.business.resolve_evidence(
        EvidenceOutcome(event_id, EvidenceStatus.UNAVAILABLE, None, None, None, 0, "x")
    )
    assert context.persistence.drain(timeout_s=2)
    decision = pipeline.business.decisions[-1]
    assert (decision.event_id, decision.decision, decision.review_reason) == (
        event_id,
        "review_required",
        "SOURCE_INTERRUPTED",
    )
    assert pipeline.analyzer.active_event().state.value == "idle"


def test_end_of_file_during_an_event_sends_it_to_review(harness: Harness) -> None:
    context = harness.context()
    evidence = FakeEvidence()
    pipeline = Pipeline(context, evidence=evidence)
    pipeline.feed(range(PICK_SEQUENCE + 2))

    pipeline.analyzer.notify_source_ended()
    pipeline.analyzer.process_pending()

    assert evidence.source_ended
    assert len(evidence.completed) == 1
    assert context.runtime.snapshot().event.state == "review_required"
    assert context.runtime.snapshot().event.review_reason == "SOURCE_INTERRUPTED"


def test_detection_failure_raises_alert_until_next_successful_frame(harness: Harness) -> None:
    context = harness.context()
    pipeline = Pipeline(context, detector=ScriptedDetector(fail_at=1))

    pipeline.feed(range(2))
    assert context.runtime.has_alert("DETECTION_FAILED")
    assert context.runtime.snapshot().analysis.last_applied_frame_sequence == 0

    pipeline.feed(range(2, 3))
    assert not context.runtime.has_alert("DETECTION_FAILED")


def test_new_event_after_completion_gets_a_new_identity(harness: Harness) -> None:
    path = PICK_PATH + [900] * 3 + list(range(850, 399, -50)) + [400] * 4
    path += list(range(450, 901, 50)) + [900] * 4
    context = harness.context()
    pipeline = Pipeline(context, detector=ScriptedDetector(path))

    pipeline.feed(range(len(path)))
    assert context.persistence.drain(timeout_s=2)

    first = pipeline.business.decisions[0].event_id
    second = pipeline.analyzer.active_event().event_id
    assert second is not None and second != first
