"""Regression tests for inventory mutation bugs found while wiring the orchestrator (PR #2)."""

from __future__ import annotations

import dataclasses
import json
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import load_workbook

from app.domain.crossing import ActionType, DirectionalAction
from app.domain.event_machine import EventContext, EventSnapshot, EventState
from app.domain.mutations import (
    MutationPrecondition,
    ObservedResult,
    ReviewCommand,
    ReviewCommandType,
    WorkbookMutation,
    canonical_hash,
)
from app.persistence.schema import SHEET_HEADERS, WorkbookFailure
from app.persistence.workbook import WorkbookRepository
from app.persistence.writer import WorkbookWriter
from app.services.inventory_service import InventoryService
from app.services.orchestrator import EventTiming, MasterData, build_review_mutation
from app.services.reconciliation import Decision, ReconciliationService
from tests.persistence.test_workbook import base_rows, write_workbook

NOW = datetime(2026, 10, 8, 16, 0, tzinfo=UTC)
STARTED = NOW - timedelta(seconds=9)
ENDED = NOW - timedelta(seconds=2)
SESSION = UUID(int=0x5E55)
EVENT_A = UUID(int=0xA)
EVENT_B = UUID(int=0xB)


def rows_with_two_tasks() -> dict[str, list[dict[str, object]]]:
    rows = base_rows()
    rows["Tasks"][0]["expected_quantity"] = 1
    rows["Tasks"].append(
        {**rows["Tasks"][0], "task_id": "task-2", "order_id": "order-2", "selected": False}
    )
    return rows


@pytest.fixture
def workbook(tmp_path: Path) -> tuple[WorkbookRepository, Path]:
    path = tmp_path / "warehouse.xlsx"
    write_workbook(path, rows_with_two_tasks())
    return WorkbookRepository(tmp_path), path


def master(repository: WorkbookRepository, path: Path, task_id: str = "task-1") -> MasterData:
    return MasterData.from_snapshot(repository.load_snapshot(path), task_id)


def event(
    event_id: UUID = EVENT_A,
    *,
    task_id: str = "task-1",
    quantity: int = 1,
    state: EventState = EventState.COMPLETED,
    confidence: float = 0.999,
    review_reason: str | None = None,
) -> EventSnapshot:
    actions = tuple(
        DirectionalAction(
            action_id=UUID(int=event_id.int * 100 + index),
            event_id=event_id,
            action_sequence=index,
            track_id=1,
            transition_index=index,
            action_type=ActionType.PICK,
            delta=1,
            source_timestamp_ms=1_000 * index,
            frame_sequence=4 * index,
            confidence=confidence,
            bbox_xyxy=(700.0, 400.0, 800.0, 500.0),
        )
        for index in range(1, quantity + 1)
    )
    return EventSnapshot(
        event_id=event_id,
        state=state,
        started_source_timestamp_ms=1_000,
        ended_source_timestamp_ms=6_000,
        pick_count=quantity,
        return_count=0,
        net_quantity=quantity,
        actions=actions,
        aggregate_confidence=confidence,
        integrity_flags=frozenset(),
        review_reason=review_reason,
        context=EventContext(
            task_id=task_id,
            zone_profile_id="zone-1",
            zone_version=1,
            sku_id="sku-box",
            unit="case",
            source_location_id="A-03-01",
            destination_location_id="STAGING",
            source_session_id=SESSION,
            configuration_version="cfg",
        ),
    )


def reconciliation() -> ReconciliationService:
    return ReconciliationService(InventoryService(), now=lambda: NOW)


def auto_approval(data: MasterData, *, workbook_version: int = 0) -> WorkbookMutation:
    result = reconciliation().reconcile(
        event(),
        data.task,
        data.source_inventory,
        data.destination_inventory,
        evidence_path="evidence/s/a.mp4",
        persistence_writable=True,
        workbook_version=workbook_version,
        started_at=STARTED,
        ended_at=ENDED,
    )
    assert result.decision is Decision.AUTO_APPROVED and result.mutation is not None
    return result.mutation


def manual_approval(data: MasterData, *, workbook_version: int = 0) -> WorkbookMutation:
    review_event = event(
        EVENT_B, task_id="task-2", state=EventState.REVIEW_REQUIRED, review_reason="LOW_CONFIDENCE"
    )
    result = reconciliation().review(
        review_event,
        data.task,
        data.source_inventory,
        data.destination_inventory,
        ReviewCommand(ReviewCommandType.APPROVE_OBSERVED, "demo_operator"),
        "review-b",
        workbook_version=workbook_version,
    )
    assert result.decision is Decision.MANUAL_APPROVED and result.mutation is not None
    return result.mutation


def all_rows(mutation: WorkbookMutation) -> dict[str, list[dict[str, object]]]:
    sheets: dict[str, list[dict[str, object]]] = {}
    for source in (mutation.upserts, mutation.appends):
        for sheet, rows in source.items():
            sheets.setdefault(sheet, []).extend(rows)
    return sheets


# --------------------------------------------------------------------------- bug 1


def test_inventory_preconditions_use_inventory_ids_with_version_and_quantity(workbook) -> None:
    data = master(*workbook)

    mutation = auto_approval(data)

    inventory = {
        item.entity_id: item.expected
        for item in mutation.preconditions
        if item.entity == "inventory"
    }
    assert inventory == {
        "inventory-source": {"version": 1, "quantity": 10},
        "inventory-destination": {"version": 1, "quantity": 0},
    }


# --------------------------------------------------------------------------- bug 2


def test_mutations_carry_the_workbook_version_they_were_built_against(workbook) -> None:
    data = master(*workbook)
    rejected = reconciliation().review(
        event(state=EventState.REVIEW_REQUIRED, review_reason="LOW_CONFIDENCE"),
        data.task,
        data.source_inventory,
        data.destination_inventory,
        ReviewCommand(ReviewCommandType.REJECT, "demo_operator", reason="wrong item"),
        "reject-a",
        workbook_version=7,
    )

    assert auto_approval(data, workbook_version=7).expected_workbook_version == 7
    manual = manual_approval(master(*workbook, "task-2"), workbook_version=7)
    assert manual.expected_workbook_version == 7
    assert rejected.mutation is not None and rejected.mutation.expected_workbook_version == 7


# --------------------------------------------------------------------------- bug 3


def test_every_written_value_uses_a_real_workbook_column(workbook) -> None:
    data = master(*workbook)
    zero = reconciliation().reconcile(
        event(quantity=0),
        data.task,
        data.source_inventory,
        data.destination_inventory,
        evidence_path=None,
        persistence_writable=True,
    )
    edited = reconciliation().review(
        event(EVENT_B, state=EventState.REVIEW_REQUIRED, review_reason="LOW_CONFIDENCE"),
        data.task,
        data.source_inventory,
        data.destination_inventory,
        ReviewCommand(
            ReviewCommandType.APPROVE_EDITED,
            "demo_operator",
            reason="counted again",
            final=ObservedResult("sku-box", 1, "case", "A-03-01"),
        ),
        "edit-b",
    )
    assert zero.mutation is not None and edited.mutation is not None

    for mutation in (auto_approval(data), zero.mutation, edited.mutation):
        for sheet, rows in all_rows(mutation).items():
            for row in rows:
                unknown = set(row) - set(SHEET_HEADERS[sheet])
                assert not unknown, (sheet, unknown)

    event_row = auto_approval(data).upserts["Events"][0]
    assert json.loads(str(event_row["integrity_flags_json"])) == []
    action = auto_approval(data).appends["EventActions"][0]
    assert json.loads(str(action["bbox_json"])) == [700.0, 400.0, 800.0, 500.0]
    review = edited.mutation.appends["Reviews"][0]
    assert json.loads(str(review["original_values_json"]))["quantity"] == 1
    assert json.loads(str(review["final_values_json"]))["source_location_id"] == "A-03-01"


# --------------------------------------------------------------------------- bug 4


def test_event_row_fills_session_task_zone_and_time_columns(workbook) -> None:
    row = auto_approval(master(*workbook)).upserts["Events"][0]

    assert row["session_id"] == str(SESSION)
    assert row["task_id"] == "task-1"
    assert (row["zone_profile_id"], row["zone_version"]) == ("zone-1", 1)
    assert row["started_at"] == STARTED.isoformat()
    assert row["ended_at"] == ENDED.isoformat()
    assert row["created_at"] == NOW.isoformat()


def test_manual_decision_keeps_existing_creation_time_times_and_evidence(workbook) -> None:
    row = manual_approval(master(*workbook, "task-2")).upserts["Events"][0]

    for column in ("created_at", "started_at", "ended_at", "evidence_path"):
        assert column not in row  # the upsert keeps the values stored at review time


# --------------------------------------------------------------------------- bug 5


def test_zero_quantity_event_is_completed_as_no_op_not_approved(workbook) -> None:
    repository, path = workbook
    data = master(repository, path)

    result = reconciliation().reconcile(
        event(quantity=0),
        data.task,
        data.source_inventory,
        data.destination_inventory,
        evidence_path="evidence/s/a.mp4",
        persistence_writable=True,
    )

    assert result.decision is Decision.NO_OP and result.mutation is not None
    row = result.mutation.upserts["Events"][0]
    assert (row["state"], row["decision"]) == ("completed", "no_op")
    assert "Inventory" not in result.mutation.upserts
    assert "Tasks" not in result.mutation.upserts
    WorkbookWriter(repository, path, now=lambda: NOW).commit(result.mutation)
    stored = repository.load_snapshot(path).indexes.event_by_id[str(EVENT_A)]
    assert (stored["state"], stored["decision"]) == ("completed", "no_op")


# --------------------------------------------------------------------------- bug 6


def test_blank_cell_matches_a_null_precondition(workbook) -> None:
    repository, path = workbook
    rows = rows_with_two_tasks()
    rows["Events"].append({"event_id": str(EVENT_B), "state": "review_required"})
    write_workbook(path, rows)
    upsert = {"Events": ({"event_id": str(EVENT_B), "review_reason": "LOW_CONFIDENCE"},)}
    mutation = WorkbookMutation(
        mutation_id=UUID(int=1),
        idempotency_key="blank-cell",
        request_id=None,
        expected_workbook_version=0,
        preconditions=(MutationPrecondition("event", str(EVENT_B), {"mutation_id": None}),),
        upserts=upsert,
        appends={},
        deletes={},
        canonical_effects_hash=canonical_hash(upsert),
    )

    receipt = WorkbookWriter(repository, path, now=lambda: NOW).commit(mutation)

    assert receipt.workbook_version_after == 1
    stale = dataclasses.replace(
        mutation,
        idempotency_key="still-checks-values",
        expected_workbook_version=1,
        preconditions=(MutationPrecondition("event", str(EVENT_B), {"state": "approved"}),),
    )
    with pytest.raises(WorkbookFailure, match="event"):
        WorkbookWriter(repository, path, now=lambda: NOW).commit(stale)


# --------------------------------------------------------------------------- end to end


def test_auto_and_manual_approvals_commit_to_a_real_workbook_and_reopen(workbook) -> None:
    repository, path = workbook
    writer = WorkbookWriter(repository, path, now=lambda: NOW)

    first = repository.load_snapshot(path)
    writer.commit(auto_approval(master(repository, path), workbook_version=first.workbook_version))

    second = repository.load_snapshot(path)
    review_event = event(
        EVENT_B, task_id="task-2", state=EventState.REVIEW_REQUIRED, review_reason="LOW_CONFIDENCE"
    )
    writer.commit(
        build_review_mutation(
            second.workbook_version,
            event=review_event,
            session_id=SESSION,
            review_reason="LOW_CONFIDENCE",
            evidence_path="evidence/s/b.mp4",
            timing=EventTiming(started_at=STARTED, ended_at=ENDED),
            created_at=ENDED,
        )
    )
    third = repository.load_snapshot(path)
    writer.commit(
        manual_approval(master(repository, path, "task-2"), workbook_version=third.workbook_version)
    )

    reopened = repository.load_snapshot(path)
    indexes = reopened.indexes
    assert reopened.workbook_version == 3
    source = indexes.inventory_by_location_sku[("A-03-01", "sku-box")]
    destination = indexes.inventory_by_location_sku[("STAGING", "sku-box")]
    assert (source["quantity"], source["version"]) == (8, 3)
    assert (destination["quantity"], destination["version"]) == (2, 3)
    assert indexes.task_by_id["task-1"]["status"] == "completed"
    assert indexes.task_by_id["task-1"]["completed_event_id"] == str(EVENT_A)
    assert indexes.task_by_id["task-2"]["status"] == "completed"

    auto_row = indexes.event_by_id[str(EVENT_A)]
    assert (auto_row["state"], auto_row["decision"]) == ("approved", "auto_approved")
    assert auto_row["session_id"] == str(SESSION)
    assert auto_row["created_at"] == NOW.isoformat()
    assert auto_row["evidence_path"] == "evidence/s/a.mp4"
    assert json.loads(str(auto_row["integrity_flags_json"])) == []

    manual_row = indexes.event_by_id[str(EVENT_B)]
    assert (manual_row["state"], manual_row["decision"]) == ("approved", "manual_approved")
    assert manual_row["created_at"] == ENDED.isoformat()
    assert manual_row["started_at"] == STARTED.isoformat()
    assert manual_row["evidence_path"] == "evidence/s/b.mp4"
    assert manual_row["final_quantity"] == 1
    [review] = indexes.reviews_by_event[str(EVENT_B)]
    assert json.loads(str(review["final_values_json"]))["quantity"] == 1

    workbook_file = load_workbook(path)
    actions = list(workbook_file["EventActions"].iter_rows(min_row=2, values_only=True))
    bbox_column = SHEET_HEADERS["EventActions"].index("bbox_json")
    assert all(json.loads(row[bbox_column]) == [700.0, 400.0, 800.0, 500.0] for row in actions)
    assert len(actions) == 2
