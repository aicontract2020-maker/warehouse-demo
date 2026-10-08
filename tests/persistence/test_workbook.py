from __future__ import annotations

import json
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import Workbook, load_workbook

from app.persistence.schema import SHEET_HEADERS, WorkbookFailure
from app.persistence.workbook import WorkbookRepository

TIMESTAMP = "2026-10-08T16:00:00+00:00"


def row(**values: object) -> dict[str, object]:
    return values


def base_rows() -> dict[str, list[dict[str, object]]]:
    zone_rectangle = json.dumps([[0.05, 0.1], [0.5, 0.1], [0.5, 0.85], [0.05, 0.85]])
    interaction = json.dumps([[0.2, 0.05], [0.8, 0.05], [0.8, 0.95], [0.2, 0.95]])
    exit_polygon = json.dumps([[0.5, 0.05], [0.98, 0.05], [0.98, 0.95], [0.5, 0.95]])
    return {
        "WorkbookMeta": [
            row(key="schema_version", value="1", updated_at=TIMESTAMP),
            row(
                key="workbook_id",
                value="00000000-0000-0000-0000-000000000001",
                updated_at=TIMESTAMP,
            ),
            row(key="scenario_id", value="basic-pick", updated_at=TIMESTAMP),
            row(key="created_at", value=TIMESTAMP, updated_at=TIMESTAMP),
            row(key="last_committed_at", value="", updated_at=TIMESTAMP),
            row(key="last_mutation_id", value="", updated_at=TIMESTAMP),
        ],
        "Skus": [
            row(
                sku_id="sku-box",
                description="Cooking oil case",
                unit="case",
                detector_class="box",
                reference_image_path="references/box.jpg",
                active=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
            )
        ],
        "Locations": [
            row(
                location_id="A-03-01",
                description="Pick face",
                location_type="pick_zone",
                active=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
            ),
            row(
                location_id="STAGING",
                description="Outbound staging",
                location_type="staging",
                active=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
            ),
        ],
        "Inventory": [
            row(
                inventory_id="inventory-source",
                location_id="A-03-01",
                sku_id="sku-box",
                quantity=10,
                unit="case",
                version=1,
                updated_at=TIMESTAMP,
                last_event_id="",
            ),
            row(
                inventory_id="inventory-destination",
                location_id="STAGING",
                sku_id="sku-box",
                quantity=0,
                unit="case",
                version=1,
                updated_at=TIMESTAMP,
                last_event_id="",
            ),
        ],
        "Tasks": [
            row(
                task_id="task-1",
                order_id="order-1",
                sku_id="sku-box",
                expected_quantity=2,
                unit="case",
                source_location_id="A-03-01",
                destination_location_id="STAGING",
                status="open",
                selected=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
                completed_event_id="",
            )
        ],
        "SourceProfiles": [
            row(
                source_profile_id="source-1",
                name="Replay source",
                source_type="replay",
                file_path="samples/pick.mp4",
                camera_index="",
                network_url_redacted="",
                analysis_fps=4,
                preview_fps=10,
                requested_width=1280,
                requested_height=720,
                reconnect_enabled=False,
                active=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
            )
        ],
        "ZoneProfiles": [
            row(
                zone_profile_id="zone-1",
                source_profile_id="source-1",
                name="A-03-01",
                version=1,
                source_width=1280,
                source_height=720,
                shelf_polygon_json=zone_rectangle,
                interaction_polygon_json=interaction,
                exit_polygon_json=exit_polygon,
                counting_line_json=json.dumps([[0.5, 0.05], [0.5, 0.95]]),
                shelf_side_sign=1,
                uncertainty_band_norm=0.02,
                crossing_confirm_frames=2,
                quiet_seconds=3,
                stable_seconds=2,
                hard_idle_seconds=10,
                merge_window_seconds=10,
                active=True,
                created_at=TIMESTAMP,
                updated_at=TIMESTAMP,
            )
        ],
        "Sessions": [],
        "Events": [],
        "EventActions": [],
        "Reviews": [],
        "AuditLog": [
            row(
                audit_id="audit-fixture",
                timestamp=TIMESTAMP,
                operator_id="system",
                action="fixture_created",
                entity_type="scenario",
                entity_id="basic-pick",
                request_id="",
                before_json="",
                after_json="",
                result="success",
                error_code="",
                details_json="{}",
            )
        ],
        "MutationReceipts": [],
    }


def write_workbook(path: Path, rows: dict[str, list[dict[str, object]]] | None = None) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    source_rows = rows or base_rows()
    for sheet_name, headers in SHEET_HEADERS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for source_row in source_rows[sheet_name]:
            sheet.append([source_row.get(header, "") for header in headers])
    workbook.save(path)


def test_loads_valid_schema_and_builds_business_indexes(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.xlsx"
    write_workbook(path)

    snapshot = WorkbookRepository(tmp_path).load_snapshot(path)

    assert snapshot.schema_version == 1
    assert snapshot.workbook_id == UUID("00000000-0000-0000-0000-000000000001")
    assert snapshot.scenario_id == "basic-pick"
    assert snapshot.workbook_version == 0
    assert snapshot.indexes.sku_by_id["sku-box"]["unit"] == "case"
    assert snapshot.indexes.inventory_by_location_sku[("A-03-01", "sku-box")]["quantity"] == 10
    assert snapshot.indexes.selected_task["task_id"] == "task-1"
    assert snapshot.indexes.active_source_profile["source_profile_id"] == "source-1"
    assert snapshot.indexes.active_zone_by_source["source-1"].version == 1


def test_rejects_missing_sheet_and_header_order_mismatch(tmp_path: Path) -> None:
    missing_path = tmp_path / "missing.xlsx"
    write_workbook(missing_path)
    workbook = load_workbook(missing_path)
    del workbook["Reviews"]
    workbook.save(missing_path)

    with pytest.raises(WorkbookFailure) as missing:
        WorkbookRepository(tmp_path).load_snapshot(missing_path)
    assert missing.value.code == "SHEET_MISSING"

    headers_path = tmp_path / "headers.xlsx"
    write_workbook(headers_path)
    workbook = load_workbook(headers_path)
    workbook["Inventory"].cell(1, 1).value = "wrong"
    workbook.save(headers_path)

    with pytest.raises(WorkbookFailure) as headers:
        WorkbookRepository(tmp_path).load_snapshot(headers_path)
    assert headers.value.code == "HEADER_MISMATCH"


def test_rejects_unsupported_schema_and_path_escape(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.xlsx"
    rows = base_rows()
    rows["WorkbookMeta"][0]["value"] = "2"
    write_workbook(path, rows)

    with pytest.raises(WorkbookFailure) as schema:
        WorkbookRepository(tmp_path).load_snapshot(path)
    assert schema.value.code == "SCHEMA_VERSION_UNSUPPORTED"

    with pytest.raises(WorkbookFailure) as escaped:
        WorkbookRepository(tmp_path / "data").load_snapshot(path)
    assert escaped.value.code == "PATH_NOT_ALLOWED"


def test_rejects_duplicate_keys_bad_rows_and_broken_foreign_keys(tmp_path: Path) -> None:
    duplicate_path = tmp_path / "duplicate.xlsx"
    rows = base_rows()
    rows["Inventory"].append(dict(rows["Inventory"][0], inventory_id="another"))
    write_workbook(duplicate_path, rows)
    with pytest.raises(WorkbookFailure) as duplicate:
        WorkbookRepository(tmp_path).load_snapshot(duplicate_path)
    assert duplicate.value.code == "DUPLICATE_KEY"

    bad_row_path = tmp_path / "bad-row.xlsx"
    rows = base_rows()
    rows["Inventory"][0]["quantity"] = -1
    write_workbook(bad_row_path, rows)
    with pytest.raises(WorkbookFailure) as bad_row:
        WorkbookRepository(tmp_path).load_snapshot(bad_row_path)
    assert bad_row.value.code == "ROW_VALIDATION_ERROR"

    broken_fk_path = tmp_path / "broken-fk.xlsx"
    rows = base_rows()
    rows["Tasks"][0]["sku_id"] = "missing-sku"
    write_workbook(broken_fk_path, rows)
    with pytest.raises(WorkbookFailure) as broken_fk:
        WorkbookRepository(tmp_path).load_snapshot(broken_fk_path)
    assert broken_fk.value.code == "REFERENTIAL_INTEGRITY_ERROR"


def test_recovers_events_actions_reviews_and_idempotency_receipts(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.xlsx"
    rows = base_rows()
    event_id = "00000000-0000-0000-0000-000000000010"
    rows["Events"].append(
        {header: "" for header in SHEET_HEADERS["Events"]}
        | {
            "event_id": event_id,
            "session_id": "",
            "task_id": "task-1",
            "zone_profile_id": "zone-1",
            "zone_version": 1,
            "state": "approved",
            "observed_sku_id": "sku-box",
            "observed_quantity": 2,
            "observed_unit": "case",
            "observed_source_location_id": "A-03-01",
            "decision": "auto_approved",
            "started_at": TIMESTAMP,
            "ended_at": TIMESTAMP,
            "created_at": TIMESTAMP,
        }
    )
    rows["EventActions"].append(
        {header: "" for header in SHEET_HEADERS["EventActions"]}
        | {
            "action_id": "action-1",
            "event_id": event_id,
            "action_sequence": 1,
            "track_id": 1,
            "transition_index": 1,
            "action_type": "pick",
            "delta": 1,
            "source_timestamp_ms": 1000,
            "frame_sequence": 4,
            "confidence": 0.999,
            "bbox_json": "[1,2,3,4]",
            "created_at": TIMESTAMP,
        }
    )
    rows["Reviews"].append(
        {header: "" for header in SHEET_HEADERS["Reviews"]}
        | {
            "review_id": "review-1",
            "event_id": event_id,
            "decision": "approve_observed",
            "operator_id": "demo_operator",
            "original_values_json": "{}",
            "final_values_json": "{}",
            "reason": "",
            "created_at": TIMESTAMP,
        }
    )
    rows["MutationReceipts"].append(
        {
            "mutation_id": "00000000-0000-0000-0000-000000000020",
            "idempotency_key": event_id,
            "event_id": event_id,
            "mutation_hash": "a" * 64,
            "committed_at": TIMESTAMP,
            "workbook_version_before": 0,
            "workbook_version_after": 1,
            "result_json": "{}",
        }
    )
    write_workbook(path, rows)

    snapshot = WorkbookRepository(tmp_path).load_snapshot(path)

    assert snapshot.workbook_version == 1
    assert snapshot.indexes.event_by_id[event_id]["decision"] == "auto_approved"
    assert snapshot.indexes.actions_by_event[event_id][0]["action_sequence"] == 1
    assert snapshot.indexes.reviews_by_event[event_id][0]["decision"] == "approve_observed"
    assert snapshot.indexes.receipt_by_idempotency_key[event_id]["mutation_hash"] == "a" * 64
