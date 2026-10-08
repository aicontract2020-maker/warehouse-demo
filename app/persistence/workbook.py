from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.domain.zones import ZoneProfile
from app.persistence.schema import (
    PRIMARY_KEYS,
    REQUIRED_METADATA_KEYS,
    SHEET_HEADERS,
    ValidationReport,
    WorkbookFailure,
)

Row = dict[str, Any]


@dataclass(frozen=True, slots=True)
class WorkbookIndexes:
    sku_by_id: dict[str, Row]
    location_by_id: dict[str, Row]
    inventory_by_location_sku: dict[tuple[str, str], Row]
    task_by_id: dict[str, Row]
    selected_task: Row | None
    source_profile_by_id: dict[str, Row]
    active_source_profile: Row | None
    zone_by_source_version: dict[tuple[str, int], ZoneProfile]
    active_zone_by_source: dict[str, ZoneProfile]
    session_by_id: dict[str, Row]
    event_by_id: dict[str, Row]
    actions_by_event: dict[str, tuple[Row, ...]]
    reviews_by_event: dict[str, tuple[Row, ...]]
    receipt_by_idempotency_key: dict[str, Row]


@dataclass(frozen=True, slots=True)
class WorkbookSnapshot:
    path: Path
    schema_version: int
    workbook_id: UUID
    scenario_id: str
    workbook_version: int
    rows: dict[str, tuple[Row, ...]]
    indexes: WorkbookIndexes


class WorkbookRepository:
    def __init__(self, data_dir: Path) -> None:
        self._data_dir = data_dir.expanduser().resolve()

    def load_snapshot(self, path: Path) -> WorkbookSnapshot:
        resolved = path.expanduser().resolve()
        if not resolved.is_relative_to(self._data_dir):
            raise WorkbookFailure("PATH_NOT_ALLOWED", "workbook must be inside data directory")
        if not resolved.is_file():
            raise WorkbookFailure("WORKBOOK_NOT_FOUND", str(resolved))

        try:
            workbook = load_workbook(resolved, data_only=False, read_only=False)
        except (InvalidFileException, OSError, ValueError, KeyError) as error:
            raise WorkbookFailure("WORKBOOK_CORRUPT", str(error)) from error

        self._validate_structure(workbook.sheetnames, workbook)
        rows = {
            sheet_name: self._read_rows(workbook[sheet_name], headers)
            for sheet_name, headers in SHEET_HEADERS.items()
        }
        report = self.validate_rows(rows)
        if not report.valid:
            code, _, detail = report.errors[0].partition(":")
            raise WorkbookFailure(code, detail.strip())

        metadata = {str(item["key"]): str(item["value"] or "") for item in rows["WorkbookMeta"]}
        indexes = self._build_indexes(rows)
        receipt_versions = [
            self._integer(item["workbook_version_after"], "receipt workbook version")
            for item in rows["MutationReceipts"]
        ]
        return WorkbookSnapshot(
            path=resolved,
            schema_version=int(metadata["schema_version"]),
            workbook_id=UUID(metadata["workbook_id"]),
            scenario_id=metadata["scenario_id"],
            workbook_version=max(receipt_versions, default=0),
            rows=rows,
            indexes=indexes,
        )

    def validate_snapshot(self, snapshot: WorkbookSnapshot) -> ValidationReport:
        return self.validate_rows(snapshot.rows)

    def validate_rows(self, rows: dict[str, tuple[Row, ...]]) -> ValidationReport:
        errors: list[str] = []
        metadata = {
            str(item.get("key", "")): str(item.get("value") or "") for item in rows["WorkbookMeta"]
        }
        missing_metadata = REQUIRED_METADATA_KEYS - metadata.keys()
        if missing_metadata:
            errors.append(f"ROW_VALIDATION_ERROR: missing metadata {sorted(missing_metadata)}")
        if metadata.get("schema_version") != "1":
            errors.append("SCHEMA_VERSION_UNSUPPORTED: expected schema version 1")
        try:
            UUID(metadata.get("workbook_id", ""))
        except ValueError:
            errors.append("ROW_VALIDATION_ERROR: workbook_id must be a UUID")

        for sheet_name, primary_key in PRIMARY_KEYS.items():
            seen: set[str] = set()
            for item in rows[sheet_name]:
                key = str(item.get(primary_key) or "")
                if not key:
                    errors.append(f"ROW_VALIDATION_ERROR: {sheet_name}.{primary_key} is required")
                elif key in seen:
                    errors.append(f"DUPLICATE_KEY: duplicate {sheet_name}.{primary_key}={key}")
                seen.add(key)

        sku_by_id = {str(item["sku_id"]): item for item in rows["Skus"]}
        location_by_id = {str(item["location_id"]): item for item in rows["Locations"]}
        source_by_id = {str(item["source_profile_id"]): item for item in rows["SourceProfiles"]}
        inventory_keys: set[tuple[str, str]] = set()
        for item in rows["Inventory"]:
            key = (str(item["location_id"]), str(item["sku_id"]))
            if key in inventory_keys:
                errors.append(f"DUPLICATE_KEY: duplicate Inventory location/SKU {key}")
            inventory_keys.add(key)
            if key[0] not in location_by_id or key[1] not in sku_by_id:
                errors.append("REFERENTIAL_INTEGRITY_ERROR: inventory location or SKU missing")
            try:
                quantity = self._integer(item["quantity"], "inventory quantity")
                version = self._integer(item["version"], "inventory version")
                if quantity < 0 or version < 1:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append("ROW_VALIDATION_ERROR: inventory quantity/version invalid")
            sku = sku_by_id.get(key[1])
            if sku and item["unit"] != sku["unit"]:
                errors.append("REFERENTIAL_INTEGRITY_ERROR: inventory unit differs from SKU")

        selected_tasks = [item for item in rows["Tasks"] if bool(item["selected"])]
        if len(selected_tasks) > 1:
            errors.append("ROW_VALIDATION_ERROR: more than one task selected")
        for item in rows["Tasks"]:
            if (
                str(item["sku_id"]) not in sku_by_id
                or str(item["source_location_id"]) not in location_by_id
                or str(item["destination_location_id"]) not in location_by_id
            ):
                errors.append("REFERENTIAL_INTEGRITY_ERROR: task reference missing")
            try:
                if self._integer(item["expected_quantity"], "task quantity") <= 0:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append("ROW_VALIDATION_ERROR: task quantity invalid")

        active_sources = [item for item in rows["SourceProfiles"] if bool(item["active"])]
        if len(active_sources) > 1:
            errors.append("ROW_VALIDATION_ERROR: more than one source profile active")
        for item in rows["SourceProfiles"]:
            self._validate_relative_path(item.get("file_path"), "source file", errors)
            if item["source_type"] not in {"file", "camera", "network", "replay"}:
                errors.append("ROW_VALIDATION_ERROR: source type invalid")
            if item["analysis_fps"] not in {3, 4, 5}:
                errors.append("ROW_VALIDATION_ERROR: source analysis_fps invalid")

        active_zone_sources: set[str] = set()
        for item in rows["ZoneProfiles"]:
            source_id = str(item["source_profile_id"])
            if source_id not in source_by_id:
                errors.append("REFERENTIAL_INTEGRITY_ERROR: zone source profile missing")
                continue
            if bool(item["active"]) and source_id in active_zone_sources:
                errors.append("ROW_VALIDATION_ERROR: multiple active zones for source")
            if bool(item["active"]):
                active_zone_sources.add(source_id)
            try:
                self._zone_from_row(item)
            except (ValidationError, ValueError, TypeError, json.JSONDecodeError) as error:
                errors.append(f"ROW_VALIDATION_ERROR: invalid zone profile: {error}")

        event_ids = {str(item["event_id"]) for item in rows["Events"]}
        for item in rows["EventActions"]:
            if str(item["event_id"]) not in event_ids:
                errors.append("REFERENTIAL_INTEGRITY_ERROR: action event missing")
        for item in rows["Reviews"]:
            if str(item["event_id"]) not in event_ids:
                errors.append("REFERENTIAL_INTEGRITY_ERROR: review event missing")
        receipt_keys: set[str] = set()
        for item in rows["MutationReceipts"]:
            key = str(item["idempotency_key"])
            if key in receipt_keys:
                errors.append("DUPLICATE_KEY: duplicate receipt idempotency key")
            receipt_keys.add(key)

        return ValidationReport(tuple(errors))

    @staticmethod
    def _validate_structure(sheetnames: list[str], workbook: Any) -> None:
        for sheet_name, expected_headers in SHEET_HEADERS.items():
            if sheet_name not in sheetnames:
                raise WorkbookFailure("SHEET_MISSING", sheet_name)
            actual = tuple(
                workbook[sheet_name].cell(1, index).value
                for index in range(1, len(expected_headers) + 1)
            )
            if actual != expected_headers or workbook[sheet_name].max_column != len(
                expected_headers
            ):
                raise WorkbookFailure("HEADER_MISMATCH", sheet_name)

    @staticmethod
    def _read_rows(sheet: Any, headers: tuple[str, ...]) -> tuple[Row, ...]:
        result: list[Row] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in values):
                continue
            result.append(
                {
                    header: "" if value is None else value
                    for header, value in zip(headers, values, strict=True)
                }
            )
        return tuple(result)

    def _build_indexes(self, rows: dict[str, tuple[Row, ...]]) -> WorkbookIndexes:
        zones = [self._zone_from_row(item) for item in rows["ZoneProfiles"]]
        actions: dict[str, list[Row]] = {}
        for item in rows["EventActions"]:
            actions.setdefault(str(item["event_id"]), []).append(item)
        reviews: dict[str, list[Row]] = {}
        for item in rows["Reviews"]:
            reviews.setdefault(str(item["event_id"]), []).append(item)
        return WorkbookIndexes(
            sku_by_id={str(item["sku_id"]): item for item in rows["Skus"]},
            location_by_id={str(item["location_id"]): item for item in rows["Locations"]},
            inventory_by_location_sku={
                (str(item["location_id"]), str(item["sku_id"])): item for item in rows["Inventory"]
            },
            task_by_id={str(item["task_id"]): item for item in rows["Tasks"]},
            selected_task=next((item for item in rows["Tasks"] if bool(item["selected"])), None),
            source_profile_by_id={
                str(item["source_profile_id"]): item for item in rows["SourceProfiles"]
            },
            active_source_profile=next(
                (item for item in rows["SourceProfiles"] if bool(item["active"])), None
            ),
            zone_by_source_version={(item.source_profile_id, item.version): item for item in zones},
            active_zone_by_source={item.source_profile_id: item for item in zones if item.active},
            session_by_id={str(item["session_id"]): item for item in rows["Sessions"]},
            event_by_id={str(item["event_id"]): item for item in rows["Events"]},
            actions_by_event={
                key: tuple(sorted(value, key=lambda item: int(item["action_sequence"])))
                for key, value in actions.items()
            },
            reviews_by_event={
                key: tuple(sorted(value, key=lambda item: str(item["created_at"])))
                for key, value in reviews.items()
            },
            receipt_by_idempotency_key={
                str(item["idempotency_key"]): item for item in rows["MutationReceipts"]
            },
        )

    @staticmethod
    def _zone_from_row(item: Row) -> ZoneProfile:
        return ZoneProfile(
            zone_profile_id=str(item["zone_profile_id"]),
            source_profile_id=str(item["source_profile_id"]),
            name=str(item["name"]),
            version=int(item["version"]),
            source_width=int(item["source_width"]),
            source_height=int(item["source_height"]),
            shelf_polygon=json.loads(str(item["shelf_polygon_json"])),
            interaction_polygon=json.loads(str(item["interaction_polygon_json"])),
            exit_polygon=json.loads(str(item["exit_polygon_json"])),
            counting_line=json.loads(str(item["counting_line_json"])),
            shelf_side_sign=int(item["shelf_side_sign"]),
            uncertainty_band_norm=float(item["uncertainty_band_norm"]),
            crossing_confirm_frames=int(item["crossing_confirm_frames"]),
            quiet_seconds=float(item["quiet_seconds"]),
            stable_seconds=float(item["stable_seconds"]),
            hard_idle_seconds=float(item["hard_idle_seconds"]),
            merge_window_seconds=float(item["merge_window_seconds"]),
            active=bool(item["active"]),
        )

    @staticmethod
    def _integer(value: object, label: str) -> int:
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError(f"{label} must be an integer")
        return value

    @staticmethod
    def _validate_relative_path(value: object, label: str, errors: list[str]) -> None:
        if value in {None, ""}:
            return
        path = Path(str(value))
        if path.is_absolute() or ".." in path.parts:
            errors.append(f"ROW_VALIDATION_ERROR: {label} path escapes data directory")
